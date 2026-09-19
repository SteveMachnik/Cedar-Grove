# -*- coding: utf-8 -*-
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

exec(open('build.py').read())

GREEN   = "375D47"
DARK    = "253F30"
CREAM   = "FFF4D6"
TINT    = "F7EAC4"
RUST    = "A16A4A"
SAGE    = "8DA57C"
BLUE    = "C2E0ED"
WHITE   = "FFFFFF"

H  = Font(name="Georgia", size=11, bold=True, color=CREAM)
TT = Font(name="Georgia", size=18, bold=True, color=DARK)
SB = Font(name="Georgia", size=12, bold=True, color=DARK)
BD = Font(name="Calibri", size=10)
BDB= Font(name="Calibri", size=10, bold=True)
IT = Font(name="Calibri", size=10, italic=True, color="555555")

fill_head = PatternFill("solid", fgColor=GREEN)
fill_tint = PatternFill("solid", fgColor=TINT)
fill_cream= PatternFill("solid", fgColor=CREAM)

thin = Side(style="thin", color="D8CBA6")
box  = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
top  = Alignment(vertical="top")
ctr  = Alignment(horizontal="center", vertical="top")

PILLARS = [
 (RH, "Rhythm of the Day",
  "The ordinary shape of a day: morning gratitude, main lesson, outdoor time, handwork, home at 3:30.",
  "Parents want to picture their child inside your day. This is the pillar that does that."),
 (NA, "Nature & Season",
  "The land itself — 700 acres of farm and forest — and how it changes through the year.",
  "Your single biggest differentiator, and the easiest content to shoot on any given day."),
 (FE, "Festivals & Gatherings",
  "The eight festivals, the build-up to each, and the recap afterwards.",
  "Your highest-engagement posts. Festivals are visual, communal, and dated — plan around them."),
 (FA, "Faces of Cedar Grove",
  "Teachers, families, testimonials, and the people who make the school.",
  "Trust is personal at this size. Faces convert better than facilities."),
 (WH, "Why We Do This",
  "The reasoning behind a choice — mixed-age classes, no screens, two hours outside, mud.",
  "Answers the quiet objection a prospective parent has but won't email you about."),
 (CS, "Come and See",
  "Tours, open houses, enrolment windows, programme explainers.",
  "The ask. Kept to roughly 1 in 8 posts so the feed never feels like advertising."),
]

wb = Workbook()

# ---------------------------------------------------------------- READ ME
ws = wb.active
ws.title = "Read Me"
ws.sheet_view.showGridLines = False
ws["B2"] = "Cedar Grove School — Content Calendar 2026–2027"; ws["B2"].font = TT
ws["B3"] = "Instagram + Facebook · roughly 2–3 posts a week · built around the school calendar"; ws["B3"].font = IT

r = 5
ws.cell(r,2,"How to use this").font = SB; r += 1
for line in [
 "The Calendar tab is the working sheet. One row per post, in date order, from the first day of school to the last.",
 "Each row gives you the date, where it goes, which pillar it serves, the idea, a caption starter, and the photo you need.",
 "Caption starters are openings, not finished captions — the last two lines should be yours.",
 "Set the Status column to Drafted / Scheduled / Posted as you go. The counts on this page update themselves.",
 "The Photo column is the important one. Read a fortnight ahead so you shoot what you'll need before the moment passes.",
]:
    c = ws.cell(r,2,"•  " + line); c.font = BD; r += 1
r += 1

ws.cell(r,2,"The six pillars").font = SB; r += 1
hdr = ["Pillar","What it covers","Why it earns its place","Posts","% of year"]
for i,h in enumerate(hdr):
    c = ws.cell(r,2+i,h); c.font = H; c.fill = fill_head; c.border = box; c.alignment = ctr
pr = r; r += 1
first, last = 4, 4 + len(P) - 1
for name,_lbl,what,why in PILLARS:
    ws.cell(r,2,name).font = BDB
    ws.cell(r,3,what).font = BD; ws.cell(r,3).alignment = wrap
    ws.cell(r,4,why).font = BD;  ws.cell(r,4).alignment = wrap
    ws.cell(r,5,'=COUNTIF(Calendar!$E$%d:$E$%d,$B%d)' % (first,last,r)).font = BD
    ws.cell(r,6,'=IFERROR($E%d/SUM($E$%d:$E$%d),0)' % (r,pr+1,pr+len(PILLARS))).font = BD
    ws.cell(r,6).number_format = "0%"
    for cc in range(2,7):
        ws.cell(r,cc).border = box
        if cc in (3,4): ws.cell(r,cc).alignment = wrap
        else: ws.cell(r,cc).alignment = top
    r += 1
tr = r
ws.cell(r,2,"Total").font = BDB
ws.cell(r,5,'=SUM($E$%d:$E$%d)' % (pr+1,r-1)).font = BDB
ws.cell(r,6,'=IFERROR($E%d/$E%d,0)' % (r,r)).font = BDB; ws.cell(r,6).number_format = "0%"
for cc in range(2,7): ws.cell(r,cc).border = box; ws.cell(r,cc).fill = fill_tint
r += 2

ws.cell(r,2,"Where it goes").font = SB; r += 1
for i,h in enumerate(["Channel","Posts"]):
    c = ws.cell(r,2+i,h); c.font = H; c.fill = fill_head; c.border = box; c.alignment = ctr
r += 1
for ch in ["Instagram + Facebook","Instagram","Facebook"]:
    ws.cell(r,2,ch).font = BD
    ws.cell(r,3,'=COUNTIF(Calendar!$D$%d:$D$%d,$B%d)' % (first,last,r)).font = BD
    for cc in (2,3): ws.cell(r,cc).border = box
    r += 1
r += 1

ws.cell(r,2,"Progress").font = SB; r += 1
for i,h in enumerate(["Status","Posts"]):
    c = ws.cell(r,2+i,h); c.font = H; c.fill = fill_head; c.border = box; c.alignment = ctr
r += 1
for st in ["Not started","Drafted","Scheduled","Posted"]:
    ws.cell(r,2,st).font = BD
    ws.cell(r,3,'=COUNTIF(Calendar!$J$%d:$J$%d,$B%d)' % (first,last,r)).font = BD
    for cc in (2,3): ws.cell(r,cc).border = box
    r += 1
r += 2

ws.cell(r,2,"Two things to fix in the school calendar PDF").font = SB; r += 1
for line in [
 'The legend reads "First day of school: Monday, September 8th, 2026." September 8th 2026 is a Tuesday.',
 'The legend reads "Last day of school: Friday, June 18th, 2026." That should be 2027. (June 18th 2027 is indeed a Friday.)',
 "Every other date in the PDF checks out — no festival falls on a weekend.",
]:
    c = ws.cell(r,2,"•  " + line); c.font = BD; r += 1
r += 1
ws.cell(r,2,"One post needs a person, not a plan").font = SB; r += 1
c = ws.cell(r,2,"September 30th is the National Day for Truth and Reconciliation. That post is flagged HOLD FOR REVIEW in the calendar and should be written or approved by Kerri (Kaya'tahente) Smart. Nothing generic.")
c.font = BD; c.alignment = wrap
ws.merge_cells(start_row=r,start_column=2,end_row=r+1,end_column=6)

for col,w in zip("ABCDEF",[3,30,52,52,10,10]): ws.column_dimensions[col].width = w

# ---------------------------------------------------------------- CALENDAR
cs = wb.create_sheet("Calendar")
cs.sheet_view.showGridLines = False
cs["A1"] = "Content Calendar — Cedar Grove School, 2026–2027"; cs["A1"].font = TT
cs["A2"] = "One row per post. Caption starters are openings — finish them in your own voice."; cs["A2"].font = IT

cols = ["Date","Day","Week of","Channel","Pillar","Post idea","Caption starter","Photo / asset needed","School tie-in","Status"]
widths = [11,6,11,20,22,30,58,34,26,13]
for i,(h,w) in enumerate(zip(cols,widths),start=1):
    c = cs.cell(3,i,h); c.font = H; c.fill = fill_head; c.border = box
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cs.column_dimensions[get_column_letter(i)].width = w
cs.row_dimensions[3].height = 26

DAYS = ["Mon","Tue","Wed","Thu","Fri","Sat","Sun"]
row = 4
prev_month = None
for d_iso, ch, pil, idea, cap, photo, tie in P:
    d = datetime.date.fromisoformat(d_iso)
    monday = d - datetime.timedelta(days=d.weekday())
    vals = [d, DAYS[d.weekday()], monday, ch, pil, idea, cap, photo, tie, "Not started"]
    for i,v in enumerate(vals, start=1):
        c = cs.cell(row,i,v)
        c.font = BD; c.border = box
        c.alignment = wrap if i in (6,7,8,9) else (ctr if i in (2,4,5,10) else top)
    cs.cell(row,1).number_format = "d mmm yyyy"
    cs.cell(row,3).number_format = "d mmm"
    if d.month != prev_month:
        for i in range(1,11): cs.cell(row,i).fill = fill_tint
        prev_month = d.month
    cs.row_dimensions[row].height = 44
    row += 1

last_row = row - 1
cs.freeze_panes = "A4"
cs.auto_filter.ref = "A3:J%d" % last_row

dv_status = DataValidation(type="list", formula1='"Not started,Drafted,Scheduled,Posted"', allow_blank=True)
cs.add_data_validation(dv_status); dv_status.add("J4:J%d" % last_row)
dv_ch = DataValidation(type="list", formula1='"Instagram + Facebook,Instagram,Facebook"', allow_blank=True)
cs.add_data_validation(dv_ch); dv_ch.add("D4:D%d" % last_row)
dv_pl = DataValidation(type="list", formula1='"%s"' % ",".join(p[0] for p in PILLARS), allow_blank=True)
cs.add_data_validation(dv_pl); dv_pl.add("E4:E%d" % last_row)

rng = "A4:J%d" % last_row
cs.conditional_formatting.add(rng, FormulaRule(formula=['$J4="Posted"'], fill=PatternFill("solid", fgColor="E8EFE4")))
cs.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("HOLD FOR REVIEW",$G4))'], fill=PatternFill("solid", fgColor="F6DCCB"), font=Font(name="Calibri", size=10, bold=True, color="8A3B1A")))

# ---------------------------------------------------------------- SCHOOL DATES
sd = wb.create_sheet("School Dates")
sd.sheet_view.showGridLines = False
sd["A1"] = "School Calendar 2026–2027 — verified"; sd["A1"].font = TT
sd["A2"] = "Every date below was checked against the actual day of the week."; sd["A2"].font = IT

EV = [
("2026-09-08","First day of school","Term","Yes — first-day post"),
("2026-09-29","Festival of Courage","Festival","Yes — build-up, day-of, recap"),
("2026-09-30","National Day for Truth and Reconciliation","Observance","Yes — teacher-led, see note"),
("2026-10-09","Harvest Festival (families welcome)","Festival","Yes — invitation, day-of, recap"),
("2026-10-12","Thanksgiving — no school","Holiday","No"),
("2026-10-26","PA Day / Parent-Teacher Conferences","PA Day","Light — reminder only"),
("2026-11-06","Lantern Festival","Festival","Yes — build-up, day-of, recap"),
("2026-11-11","Remembrance Day","Observance","Yes — simple and restrained"),
("2026-11-20","PA Day","PA Day","Light — reminder only"),
("2026-12-18","Spiral Walk (families welcome)","Festival","Yes — build-up, day-of, recap"),
("2026-12-21","Winter Holidays begin","Holiday","Yes — sign-off post"),
("2027-01-05","First day back","Term","Yes"),
("2027-01-29","PA Day","PA Day","Light — reminder only"),
("2027-02-05","Mid-Winter Festival","Festival","Yes — build-up, day-of, recap"),
("2027-02-12","PA Day / Parent-Teacher Conferences","PA Day","Light — reminder only"),
("2027-02-15","Family Day — no school","Holiday","Light"),
("2027-03-08","Spring Break begins","Holiday","Yes — sign-off"),
("2027-03-19","Spring Break ends","Holiday","Yes — return reminder"),
("2027-03-26","Good Friday — no school","Holiday","No"),
("2027-03-29","Easter Monday — no school","Holiday","No"),
("2027-04-02","Seed Festival (families welcome)","Festival","Yes — build-up, day-of, recap"),
("2027-04-23","PA Day","PA Day","Light — reminder only"),
("2027-05-14","Mayfair","Festival","Yes — build-up, day-of, recap"),
("2027-05-24","Victoria Day — no school","Holiday","No"),
("2027-05-31","PA Day","PA Day","Light — reminder only"),
("2027-06-18","Strawberry Festival & last day of school","Festival","Yes — build-up, day-of, farewell"),
]
for i,(h,w) in enumerate(zip(["Date","Day","Event","Type","Content?"],[13,7,44,13,32]),start=1):
    c = sd.cell(4,i,h); c.font = H; c.fill = fill_head; c.border = box; c.alignment = ctr
    sd.column_dimensions[get_column_letter(i)].width = w
rr = 5
for iso, name, typ, use in EV:
    d = datetime.date.fromisoformat(iso)
    for i,v in enumerate([d, DAYS[d.weekday()], name, typ, use], start=1):
        c = sd.cell(rr,i,v); c.font = BD; c.border = box
        c.alignment = ctr if i in (2,4) else top
    sd.cell(rr,1).number_format = "d mmm yyyy"
    if typ == "Festival":
        for i in range(1,6): sd.cell(rr,i).fill = fill_tint
    rr += 1
sd.freeze_panes = "A5"

# ---------------------------------------------------------------- SHOT LIST
sl = wb.create_sheet("Shot List")
sl.sheet_view.showGridLines = False
sl["A1"] = "What to photograph, and roughly when"; sl["A1"].font = TT
sl["A2"] = "Shoot ahead. Most of these can only be got in one season."; sl["A2"].font = IT

SHOTS = [
("September","Children walking down the trail, shot from behind","The single most Cedar Grove image you own. Get it in every season."),
("September","Morning Gratitude — the whole school in circle","Anchors the 'rhythm' pillar all year."),
("September","Each teacher, portrait plus one candid at work","Reused constantly. Worth an hour of someone's time."),
("September","First maples turning along the trail","Two-week window."),
("October","Harvest Festival — apple pressing, the table, families","Your first family-facing festival."),
("October","Boots, mittens and wool laid out flat","Used again in December. Shoot once."),
("November","Lanterns being made, then lit at dusk","The dusk shots are hard; take many."),
("November","Main lesson books, open, good light","Sells the academics better than any sentence."),
("December","Spiral of evergreen, one candle, one child","Low light — steady hands, high ISO."),
("December","Baking day: dough, flour, small hands","Warm, tactile, always performs."),
("January","Winter forest, tracks in snow, bare trees","Needed for the whole admissions push."),
("January","One full day, morning to home time, as a sequence","The 'what a day looks like' carousel."),
("February","Mid-Winter Festival","Interior, warm light."),
("February","A teacher with a small group of children","Proves the 8:1 ratio visually."),
("March","Sap running, buckets, boiling","Roughly a three-week window."),
("March","Seed trays and soil being filled","Ties to the native tree programme."),
("April","Mud — boots, puddles, delighted faces","Best content of the year. Do not skip it."),
("April","Maypole practice","The rehearsal shots are better than the day itself."),
("May","Mayfair — ribbons, movement, sun","Shoot wide and close."),
("May","The garden, planted and tended by children","Late May is the peak."),
("June","Strawberry Festival and the last day","Bookends the September trail shot."),
("Any time","Wide landscape of the farm and forest","Have four of these, one per season, always ready."),
]
for i,(h,w) in enumerate(zip(["When","Shot","Why it matters"],[14,52,58]),start=1):
    c = sl.cell(4,i,h); c.font = H; c.fill = fill_head; c.border = box; c.alignment = ctr
    sl.column_dimensions[get_column_letter(i)].width = w
rr = 5
prev = None
for when, shot, why in SHOTS:
    for i,v in enumerate([when, shot, why], start=1):
        c = sl.cell(rr,i,v); c.font = BD; c.border = box; c.alignment = wrap
    if when != prev:
        for i in range(1,4): sl.cell(rr,i).fill = fill_tint
        prev = when
    sl.row_dimensions[rr].height = 30
    rr += 1
sl.freeze_panes = "A5"

out = "/sessions/sleepy-upbeat-ride/mnt/Cedar Grove/marketing/Cedar Grove School - Content Calendar 2026-2027.xlsx"
wb.save(out)
print("saved", out, "rows:", last_row-3)
